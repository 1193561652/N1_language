from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
INPUT = ROOT / "01 Lessons" / "Lesson001 Business Japanese - Meeting - Safe Furigana.docx"
OUTPUT = ROOT / "01 Lessons" / "Lesson001 商务日语 - 会议 - 安全注音版.docx"


REPLACEMENTS = {
    "Lesson001 Business Japanese - Meeting": "Lesson001 商务日语 - 会议",
    "Safe Furigana Version": "安全注音版",
    "Status:": "状态：",
    "System:": "系统：",
    "Scene:": "场景：",
    "Goal:": "目标：",
    "Business Japanese / Meeting": "商务日语 / 会议",
    "Business Japanese": "商务日语",
    "Meeting": "会议",
    "① Scene": "① 场景介绍",
    "② Core Vocabulary": "② 核心词汇",
    "③ Collocations": "③ 固定搭配",
    "④ Sentence Pattern": "④ 句型表达",
    "⑤ Dialogue": "⑤ 对话",
    "⑥ Japanese Thinking": "⑥ 日本人的表达逻辑",
    "⑦ Grammar": "⑦ 文法",
    "⑧ JLPT Analysis": "⑧ JLPT 真题分析",
    "⑨ Practice": "⑨ 练习",
    "⑩ Review": "⑩ 复习",
    "Vocabulary Card": "词汇卡",
    "Function:": "功能：",
    "Formality:": "正式程度：",
    "Register:": "语域：",
    "Examples:": "例句：",
    "Thinking:": "表达逻辑：",
    "Practice:": "练习：",
    "Answer:": "答案：",
    "Answers:": "答案：",
    "Model answers:": "参考答案：",
    "Model:": "参考写法：",
    "Script:": "听力脚本：",
    "Question:": "问题：",
    "Questions:": "问题：",
    "Key expressions:": "关键表达：",
    "Source:": "来源：",
    "Source status:": "来源状态：",
    "Vocabulary:": "词汇：",
    "Collocations:": "固定搭配：",
    "Grammar:": "文法：",
    "Scene tags:": "场景标签：",
    "JLPT Connection:": "JLPT 关联：",
    "Meeting usage:": "会议场景用法：",
    "Meaning:": "意思：",
    "Text:": "阅读文本：",
    "Status": "状态",
    "Source status": "来源状态",
    "Function | Collocation | Meaning | Scene": "功能 | 固定搭配 | 意思 | 场景",
    "Lesson001 Database Items": "Lesson001 数据库条目",
    "Confirmed official source": "已确认的官方来源",
    "Uploaded past papers": "已上传真题",
    "Expected N1 links": "预期 N1 关联",
    "Collocation Practice": "固定搭配练习",
    "Sentence Building": "造句练习",
    "Instant Response": "即时应答",
    "Listening Mini Task": "听力小练习",
    "Reading Mini Task": "阅读小练习",
    "Integrated Practice": "综合练习",
    "Day 1": "第 1 天",
    "Day 2": "第 2 天",
    "Day 3": "第 3 天",
    "Week 1": "第 1 周",
    "Month 1": "第 1 个月",
    "Source-derived": "来源派生",
    "Source-informed": "来源参考",
    "Teaching rewrite": "教学改写",
}


def replace_text(text: str) -> str:
    for old, new in REPLACEMENTS.items():
        text = text.replace(old, new)
    return text


def replace_paragraph(paragraph):
    original = paragraph.text
    updated = replace_text(original)
    if updated == original:
        return
    for run in paragraph.runs:
        run.text = ""
    if paragraph.runs:
        paragraph.runs[0].text = updated
    else:
        paragraph.add_run(updated)


def main():
    doc = Document(INPUT)
    for paragraph in doc.paragraphs:
        replace_paragraph(paragraph)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    replace_paragraph(paragraph)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
