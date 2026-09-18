from pathlib import Path
import re

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "01 Lessons" / "Lesson001 Business Japanese - Meeting.md"
OUTPUT = ROOT / "01 Lessons" / "Lesson001 Business Japanese - Meeting - Safe Furigana.docx"


READINGS = {
    "日本語能力試験": "にほんごのうりょくしけん",
    "販売計画": "はんばいけいかく",
    "意見交換": "いけんこうかん",
    "関係部署": "かんけいぶしょ",
    "調査結果": "ちょうさけっか",
    "関係者": "かんけいしゃ",
    "参加者": "さんかしゃ",
    "担当者": "たんとうしゃ",
    "参考資料": "さんこうしりょう",
    "会議資料": "かいぎしりょう",
    "今日中": "きょうじゅう",
    "本日中": "ほんじつじゅう",
    "午後三時": "ごごさんじ",
    "新企画": "しんきかく",
    "会議": "かいぎ",
    "確認": "かくにん",
    "資料": "しりょう",
    "日程": "にってい",
    "調整": "ちょうせい",
    "議題": "ぎだい",
    "共有": "きょうゆう",
    "予定": "よてい",
    "内容": "ないよう",
    "連絡": "れんらく",
    "報告": "ほうこく",
    "説明": "せつめい",
    "提案": "ていあん",
    "相談": "そうだん",
    "参加": "さんか",
    "出席": "しゅっせき",
    "欠席": "けっせき",
    "変更": "へんこう",
    "決定": "けってい",
    "案内": "あんない",
    "作成": "さくせい",
    "配布": "はいふ",
    "進捗": "しんちょく",
    "情報": "じょうほう",
    "方針": "ほうしん",
    "質問": "しつもん",
    "整理": "せいり",
    "計画": "けいかく",
    "意見": "いけん",
    "事前": "じぜん",
    "承知": "しょうち",
    "添付": "てんぷ",
    "皆様": "みなさま",
    "場合": "ばあい",
    "明日": "あした",
    "今日": "きょう",
    "本日": "ほんじつ",
    "今後": "こんご",
    "来月": "らいげつ",
    "時間": "じかん",
    "都合": "つごう",
    "結論": "けつろん",
    "修正": "しゅうせい",
}

KEYS = sorted(READINGS, key=len, reverse=True)


def add_readings(text: str) -> str:
    # Only annotate lines that already contain kana, so Chinese explanations are left alone.
    if not re.search(r"[\u3040-\u30ff]", text):
        return text
    out = []
    i = 0
    while i < len(text):
        key = next((k for k in KEYS if text.startswith(k, i)), None)
        if key:
            reading = READINGS[key]
            explicit = f"{key}（{reading}）"
            if text.startswith(explicit, i):
                out.append(explicit)
                i += len(explicit)
            else:
                out.append(f"{key}（{reading}）")
                i += len(key)
        else:
            out.append(text[i])
            i += 1
    return "".join(out)


def set_font(run, size=11, bold=False):
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Microsoft YaHei")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Microsoft YaHei")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.bold = bold


def add_para(doc, text, style=None, size=11, bold=False):
    p = doc.add_paragraph(style=style)
    r = p.add_run(text)
    set_font(r, size=size, bold=bold)
    return p


def main():
    md = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)

    add_para(doc, "Lesson001 Business Japanese - Meeting", size=20, bold=True)
    add_para(doc, "Safe Furigana Version: 普通文本格式, 例如 会議（かいぎ）", size=10)

    for raw in md:
        line = raw.rstrip()
        if not line:
            continue
        if line.startswith("# "):
            continue
        if line.startswith("## "):
            add_para(doc, line[3:], style="Heading 1", size=15, bold=True)
            continue
        if line.startswith("### "):
            add_para(doc, line[4:], style="Heading 2", size=13, bold=True)
            continue
        if line.startswith("#### "):
            add_para(doc, line[5:], style="Heading 3", size=12, bold=True)
            continue
        clean = line
        if clean.startswith("- "):
            clean = "• " + clean[2:]
        clean = re.sub(r"`([^`]+)`", r"\1", clean)
        add_para(doc, add_readings(clean), size=11)

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
